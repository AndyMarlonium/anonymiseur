# -*- coding: utf-8 -*-
"""
Import/export de documents pour l'anonymiseur.

Limite assumée pour ce squelette : lors de l'export .docx, chaque
paragraphe est réécrit en un seul "run" avec le texte anonymisé.
Le style du paragraphe (titre, gras de paragraphe, alignement) est
conservé, mais une mise en forme caractère-par-caractère très fine
(un mot en gras au milieu d'une phrase) peut être simplifiée.
C'est un compromis raisonnable pour un outil d'anonymisation ; à
affiner si le rendu final doit être identique au pixel près.
"""

from __future__ import annotations
from pathlib import Path
from typing import List
import docx  # python-docx


def read_text(path: str) -> str:
    """Charge le texte d'un fichier .docx ou .txt.
    Les paragraphes .docx sont joints par des sauts de ligne."""
    p = Path(path)
    if p.suffix.lower() == ".docx":
        d = docx.Document(str(p))
        return "\n".join(par.text for par in d.paragraphs)
    else:
        return p.read_text(encoding="utf-8", errors="replace")


def write_text(path: str, anonymized_text: str, original_docx_path: str | None = None) -> None:
    """Sauvegarde le texte anonymisé.
    - Si le chemin de sortie se termine par .docx et qu'un document
      original .docx est fourni, on réutilise sa structure de
      paragraphes (styles conservés) et on remplace le texte
      paragraphe par paragraphe.
    - Sinon, écriture en texte brut ou en nouveau .docx simple."""
    out = Path(path)
    new_lines = anonymized_text.split("\n")

    if out.suffix.lower() == ".docx":
        if original_docx_path:
            d = docx.Document(original_docx_path)
            paragraphs = d.paragraphs
            # Sécurité : si le nombre de paragraphes ne correspond plus
            # (texte édité à la main dans l'aperçu), on retombe sur un
            # document neuf plutôt que de désaligner le contenu.
            if len(paragraphs) == len(new_lines):
                for par, new_text in zip(paragraphs, new_lines):
                    _replace_paragraph_text(par, new_text)
                d.save(str(out))
                return
        # Fallback : nouveau document simple
        d = docx.Document()
        for line in new_lines:
            d.add_paragraph(line)
        d.save(str(out))
    else:
        out.write_text(anonymized_text, encoding="utf-8")


def _replace_paragraph_text(paragraph, new_text: str) -> None:
    """Remplace le texte d'un paragraphe python-docx en conservant le
    style du paragraphe et la mise en forme du premier run existant."""
    if not paragraph.runs:
        paragraph.add_run(new_text)
        return
    # Garde la mise en forme du 1er run, vide les autres
    paragraph.runs[0].text = new_text
    for run in paragraph.runs[1:]:
        run.text = ""
